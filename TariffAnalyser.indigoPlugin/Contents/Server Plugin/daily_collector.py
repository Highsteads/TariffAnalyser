#! /usr/bin/env python
# -*- coding: utf-8 -*-
# Filename:    daily_collector.py
# Description: Daily energy summary collector — fetches Octopus consumption
#              and gas data, aggregates Sigenergy half-hourly data, and upserts
#              one row per calendar day into the daily_summary table.
#              Runs rolling 7-day refresh to handle Octopus data delays
#              (gas up to 2 days late, export up to 4 days late).
# Author:      CliveS & Claude Sonnet 4.6
# Date:        06-05-2026
# Version:     1.0

import base64
import json
import sqlite3
import urllib.request
import urllib.error
import urllib.parse
from datetime import datetime, date, timedelta

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_API_BASE             = "https://api.octopus.energy/v1"
_PAGE_SIZE            = 25000
_REQUEST_TIMEOUT      = 30

ELEC_STANDING_P_DAY   = 61.52    # pence/day
GAS_STANDING_P_DAY    = 29.06    # pence/day
EXPORT_RATE_P         = 12.0     # pence/kWh — Octopus Outgoing flat rate

# Octopus Go rates
GO_CHEAP_P            = 7.5      # pence/kWh 00:30-05:30
GO_PEAK_P             = 24.0     # pence/kWh all other times
GO_CHEAP_START        = "00:30"
GO_CHEAP_END          = "05:30"
GO_STANDING_P_DAY     = 53.35

# Octopus Flux rates — current published rates, region F.
# Fetched live from API where possible; these serve as fallback.
FLUX_OFFPEAK_P        = 7.01     # pence/kWh 02:00-05:00
FLUX_SHOULDER_P       = 21.00    # pence/kWh (all other hours)
FLUX_PEAK_P           = 33.00    # pence/kWh 16:00-19:00
FLUX_OFFPEAK_START    = "02:00"
FLUX_OFFPEAK_END      = "05:00"
FLUX_PEAK_START       = "16:00"
FLUX_PEAK_END         = "19:00"
FLUX_STANDING_P_DAY   = 53.35

# Gas kWh conversion: Octopus consumption API returns cubic metres.
# kWh = m3 * volume_correction (1.02264) * calorific_value (40 MJ/m3) / 3.6
GAS_KWH_PER_M3        = 11.363   # = 1.02264 * 40 / 3.6

# Tracker product for rate history backfill (region F)
TRACKER_PRODUCT       = "SILVER-25-04-11"
TRACKER_TARIFF_PREFIX = "E-1R-SILVER-25-04-11-"

SOLAR_INSTALL_DATE    = date(2026, 3, 13)   # first full day of operation


# ---------------------------------------------------------------------------
# DB schema
# ---------------------------------------------------------------------------

_CREATE_SQL = """
CREATE TABLE IF NOT EXISTS daily_summary (
    date                        TEXT PRIMARY KEY,   -- YYYY-MM-DD (UK local date)

    -- Sigenergy energy flows (from halfhourly table or File 2 import)
    pv_kwh                      REAL    DEFAULT 0.0,
    home_kwh                    REAL    DEFAULT 0.0,
    grid_import_kwh             REAL    DEFAULT 0.0,
    grid_export_kwh             REAL    DEFAULT 0.0,
    battery_charge_kwh          REAL    DEFAULT 0.0,
    battery_discharge_kwh       REAL    DEFAULT 0.0,

    -- Octopus metered electricity (may differ slightly from Sigenergy)
    octopus_import_kwh          REAL,
    octopus_export_kwh          REAL,

    -- Electricity costs — Tracker (actual tariff)
    tracker_avg_rate_p          REAL,               -- weighted avg p/kWh for the day
    elec_import_cost_gbp        REAL    DEFAULT 0.0,
    elec_export_revenue_gbp     REAL    DEFAULT 0.0,
    elec_net_cost_gbp           REAL    DEFAULT 0.0, -- import cost - export revenue
    elec_standing_charge_gbp    REAL    DEFAULT 0.6152,
    elec_total_gbp              REAL    DEFAULT 0.0, -- net cost + standing charge

    -- Savings vs no solar (Tracker rate)
    cost_without_solar_gbp      REAL    DEFAULT 0.0, -- home_kwh * tracker_rate
    savings_vs_no_solar_gbp     REAL    DEFAULT 0.0, -- avoided import + export revenue

    -- Hypothetical: Octopus Go (same consumption, TOU rates applied)
    go_import_cost_gbp          REAL,
    go_net_cost_gbp             REAL,
    go_saving_vs_tracker_gbp    REAL,               -- +ve means Go cheaper than Tracker

    -- Hypothetical: Octopus Flux (same consumption, TOU rates applied)
    flux_import_cost_gbp        REAL,
    flux_net_cost_gbp           REAL,
    flux_saving_vs_tracker_gbp  REAL,               -- +ve means Flux cheaper than Tracker

    -- Gas (from Octopus gas API)
    gas_kwh                     REAL,
    gas_unit_rate_p             REAL,
    gas_cost_gbp                REAL,
    gas_standing_charge_gbp     REAL    DEFAULT 0.2906,
    gas_total_gbp               REAL,

    -- Metadata
    sigen_source                TEXT,   -- 'halfhourly' | 'file_import' | 'none'
    octopus_elec_complete       INTEGER DEFAULT 0,  -- 1 when Octopus elec data received
    octopus_gas_complete        INTEGER DEFAULT 0,  -- 1 when Octopus gas data received
    last_updated                TEXT                -- ISO timestamp
);

CREATE INDEX IF NOT EXISTS idx_daily_summary_date ON daily_summary(date);
"""


def init_daily_summary_db(db_path):
    """Create daily_summary table and index if they do not exist."""
    con = sqlite3.connect(db_path)
    con.executescript(_CREATE_SQL)
    con.commit()
    con.close()


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def update_daily_summary(db_path, date_from, date_to, config, log_fn=None):
    """Upsert daily_summary rows for date_from..date_to (inclusive).

    config dict keys (all optional — missing keys disable that data source):
        api_key         Octopus API key (for authenticated endpoints)
        mpan            Electricity import meter point
        serial          Electricity import meter serial
        export_mpan     Electricity export meter point
        export_serial   Electricity export meter serial
        mprn            Gas meter point
        gas_serial      Gas meter serial
        region          Octopus region letter (default 'F')
        gas_unit_rate_p Gas unit rate in pence/kWh (for cost calc)
        timeseries_db   Path to SigenEnergyManager energy_timeseries.db
    """
    def _log(msg, level="INFO"):
        if log_fn:
            log_fn(f"[DailyCollector] {msg}", level=level)

    init_daily_summary_db(db_path)

    region  = config.get("region", "F")
    api_key = config.get("api_key", "")

    _log(f"Updating {date_from} to {date_to} ({(date_to - date_from).days + 1} days)")

    # ------------------------------------------------------------------ fetch
    tracker_rates  = _fetch_tracker_rates(date_from, date_to, region, _log)
    flux_rates     = _fetch_flux_rates(region, _log)
    elec_slots     = _fetch_octopus_consumption(
        api_key, config.get("mpan", ""), config.get("serial", ""),
        "electricity", date_from, date_to, _log
    )
    export_slots   = _fetch_octopus_consumption(
        api_key, config.get("export_mpan", ""), config.get("export_serial", ""),
        "export", date_from, date_to, _log
    )
    gas_daily      = _fetch_octopus_gas(
        api_key, config.get("mprn", ""), config.get("gas_serial", ""),
        date_from, date_to, _log
    )
    sigen_daily    = _aggregate_halfhourly(
        config.get("timeseries_db", ""), date_from, date_to, _log
    )
    existing_rows  = _load_existing_summary_rows(db_path, date_from, date_to)

    gas_unit_rate_p = float(config.get("gas_unit_rate_p", 6.09))

    # ------------------------------------------------------------------ build
    rows_written = 0
    cur = date_from
    while cur <= date_to:
        ds = cur.isoformat()

        # Sigenergy data — prefer halfhourly aggregation; fall back to existing row
        sg          = sigen_daily.get(ds, {})
        existing    = existing_rows.get(ds, {})
        sigen_src   = "halfhourly" if sg else ("file_import" if existing else "none")

        pv_kwh    = sg.get("pv_kwh")    or existing.get("pv_kwh")    or 0.0
        home_kwh  = sg.get("home_kwh")  or existing.get("home_kwh")  or 0.0
        imp_kwh   = sg.get("imp_kwh")   or existing.get("grid_import_kwh")  or 0.0
        exp_kwh   = sg.get("exp_kwh")   or existing.get("grid_export_kwh")  or 0.0
        bat_chg   = sg.get("bat_chg")   or existing.get("battery_charge_kwh")    or 0.0
        bat_dis   = sg.get("bat_dis")   or existing.get("battery_discharge_kwh") or 0.0

        # Octopus metered electricity for this day
        oct_imp_slots = elec_slots.get(ds, [])   # list of (time_str, kwh)
        oct_exp_slots = export_slots.get(ds, [])
        oct_imp_kwh   = sum(s[1] for s in oct_imp_slots)
        oct_exp_kwh   = sum(s[1] for s in oct_exp_slots)
        oct_elec_ok   = 1 if oct_imp_slots else 0

        # Tracker rate for this day
        tracker_p = tracker_rates.get(ds)
        if tracker_p is None and sg.get("tracker_avg_p"):
            tracker_p = sg["tracker_avg_p"]   # from halfhourly table

        # Use Octopus metered import for cost calcs where available,
        # fall back to Sigenergy figure
        billing_imp_kwh = oct_imp_kwh if oct_elec_ok else imp_kwh
        billing_exp_kwh = oct_exp_kwh if oct_exp_slots else exp_kwh

        # Electricity costs on Tracker — export revenue is always known (flat
        # rate), but import cost needs the Tracker rate. When the rate is
        # missing we write NULL for the import-derived fields rather than a
        # real-looking £0 day (a missing rate must be distinguishable from a
        # genuinely zero-cost day).
        elec_exp_rev_p = billing_exp_kwh * EXPORT_RATE_P
        if tracker_p is not None:
            elec_imp_cost_p = billing_imp_kwh * tracker_p
            elec_net_p      = elec_imp_cost_p - elec_exp_rev_p
            elec_total_p    = elec_net_p + ELEC_STANDING_P_DAY
        else:
            elec_imp_cost_p = None
            elec_net_p      = None
            elec_total_p    = None

        # Savings vs no solar — only meaningful from solar install date
        if cur >= SOLAR_INSTALL_DATE and tracker_p and home_kwh > 0:
            cost_no_solar_p    = home_kwh * tracker_p
            savings_p          = (cost_no_solar_p - elec_imp_cost_p) + elec_exp_rev_p
        elif tracker_p is not None:
            cost_no_solar_p    = billing_imp_kwh * tracker_p
            savings_p          = 0.0
        else:
            cost_no_solar_p    = None
            savings_p          = None

        # Go cost — apply TOU to Octopus half-hourly slots. "Saving vs Tracker"
        # folds each tariff's OWN standing charge into the total so the two
        # tariffs are compared like-for-like (Go/Flux standing < Tracker's, so
        # ignoring it understated the real saving).
        go_imp_p = None
        if oct_imp_slots:
            go_imp_p = _apply_tou_simple(
                oct_imp_slots, GO_CHEAP_START, GO_CHEAP_END, GO_CHEAP_P, GO_PEAK_P
            )
        go_net_p    = (go_imp_p - elec_exp_rev_p) if go_imp_p is not None else None
        go_save_p   = (
            (elec_imp_cost_p + ELEC_STANDING_P_DAY) - (go_imp_p + GO_STANDING_P_DAY)
            if (go_imp_p is not None and elec_imp_cost_p is not None) else None
        )

        # Flux cost — apply 3-band TOU
        flux_op  = flux_rates.get("offpeak_p", FLUX_OFFPEAK_P)
        flux_sh  = flux_rates.get("shoulder_p", FLUX_SHOULDER_P)
        flux_pk  = flux_rates.get("peak_p", FLUX_PEAK_P)
        flux_imp_p = None
        if oct_imp_slots:
            flux_imp_p = _apply_tou_flux(oct_imp_slots, flux_op, flux_sh, flux_pk)
        flux_net_p    = (flux_imp_p - elec_exp_rev_p) if flux_imp_p is not None else None
        flux_save_p   = (
            (elec_imp_cost_p + ELEC_STANDING_P_DAY) - (flux_imp_p + FLUX_STANDING_P_DAY)
            if (flux_imp_p is not None and elec_imp_cost_p is not None) else None
        )

        # Gas — _fetch_octopus_gas already converts m3 → kWh
        gas_kwh  = gas_daily.get(ds)
        gas_cost_p       = (gas_kwh * gas_unit_rate_p) if gas_kwh is not None else None
        gas_total_p      = (gas_cost_p + GAS_STANDING_P_DAY) if gas_cost_p is not None else None
        oct_gas_ok       = 1 if gas_kwh is not None else 0

        def _gbp(p):
            return round(p / 100.0, 4) if p is not None else None

        row = {
            "date":                     ds,
            "pv_kwh":                   round(pv_kwh, 3),
            "home_kwh":                 round(home_kwh, 3),
            "grid_import_kwh":          round(imp_kwh, 3),
            "grid_export_kwh":          round(exp_kwh, 3),
            "battery_charge_kwh":       round(bat_chg, 3),
            "battery_discharge_kwh":    round(bat_dis, 3),
            "octopus_import_kwh":       round(oct_imp_kwh, 3) if oct_elec_ok else None,
            "octopus_export_kwh":       round(oct_exp_kwh, 3) if oct_exp_slots else None,
            "tracker_avg_rate_p":       round(tracker_p, 4) if tracker_p else None,
            "elec_import_cost_gbp":     _gbp(elec_imp_cost_p),
            "elec_export_revenue_gbp":  _gbp(elec_exp_rev_p),
            "elec_net_cost_gbp":        _gbp(elec_net_p),
            "elec_standing_charge_gbp": round(ELEC_STANDING_P_DAY / 100, 4),
            "elec_total_gbp":           _gbp(elec_total_p),
            "cost_without_solar_gbp":   _gbp(cost_no_solar_p),
            "savings_vs_no_solar_gbp":  _gbp(savings_p),
            "go_import_cost_gbp":       _gbp(go_imp_p),
            "go_net_cost_gbp":          _gbp(go_net_p),
            "go_saving_vs_tracker_gbp": _gbp(go_save_p),
            "flux_import_cost_gbp":     _gbp(flux_imp_p),
            "flux_net_cost_gbp":        _gbp(flux_net_p),
            "flux_saving_vs_tracker_gbp": _gbp(flux_save_p),
            "gas_kwh":                  gas_kwh,
            "gas_unit_rate_p":          gas_unit_rate_p,
            "gas_cost_gbp":             _gbp(gas_cost_p),
            "gas_standing_charge_gbp":  round(GAS_STANDING_P_DAY / 100, 4),
            "gas_total_gbp":            _gbp(gas_total_p),
            "sigen_source":             sigen_src,
            "octopus_elec_complete":    oct_elec_ok,
            "octopus_gas_complete":     oct_gas_ok,
            "last_updated":             datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        }

        _upsert_row(db_path, row)
        rows_written += 1
        cur += timedelta(days=1)

    _log(f"Written {rows_written} daily_summary rows")
    return rows_written


def get_flux_rates(region="F", log_fn=None):
    """Return current live Flux rates for display/config purposes."""
    return _fetch_flux_rates(region, log_fn)


def get_daily_summary_rows(db_path, date_from, date_to):
    """Return list of daily_summary rows as dicts, ordered by date."""
    try:
        con  = sqlite3.connect(db_path)
        con.row_factory = sqlite3.Row
        rows = con.execute(
            "SELECT * FROM daily_summary WHERE date >= ? AND date <= ? ORDER BY date",
            (date_from.isoformat(), date_to.isoformat())
        ).fetchall()
        con.close()
        return [dict(r) for r in rows]
    except Exception:
        return []


# ---------------------------------------------------------------------------
# Sigenergy File 2 (daily XLSX) import
# ---------------------------------------------------------------------------

SIGEN_FILE2_PATH = "/Users/indigo/Documents/excel-logs-file 4/stationData-739105.xlsx"


def import_sigenergy_file2(db_path, xlsx_path=None, log_fn=None):
    """Read daily energy data from Sigenergy export XLSX and upsert into daily_summary.

    File 2 columns (daily kWh):
      Date | Solar Production Energy | Load Consumed Energy |
      Battery Charge Energy | Battery Discharge Energy |
      Grid Imported Energy | Grid Exported Energy | Revenue(GBP)

    Returns number of rows imported.
    """
    def _log(msg, level="INFO"):
        if log_fn:
            log_fn(msg, level=level)

    if xlsx_path is None:
        xlsx_path = SIGEN_FILE2_PATH

    _log(f"Importing Sigenergy File 2: {xlsx_path}")

    rows = None
    try:
        import openpyxl
        wb   = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws   = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except ImportError:
        _log("openpyxl not installed — trying xlrd", level="WARNING")
    except Exception as exc:
        if "zip" in str(exc).lower() or "not a zip" in str(exc).lower():
            _log("File is XLS format — trying xlrd", level="WARNING")
        else:
            _log(f"Cannot open {xlsx_path}: {exc}", level="ERROR")
            return 0

    if rows is None:
        try:
            import xlrd
            wb   = xlrd.open_workbook(xlsx_path)
            ws   = wb.sheet_by_index(0)
            rows = [ws.row_values(r) for r in range(ws.nrows)]
        except ImportError:
            _log("xlrd not installed — cannot read XLS file. Restart plugin to install, or re-export from Sigenergy app as XLSX.", level="ERROR")
            return 0
        except Exception as exc:
            _log(f"Cannot open {xlsx_path} with xlrd: {exc}", level="ERROR")
            return 0

    if not rows:
        _log("File is empty", level="WARNING")
        return 0

    header = [str(c).strip() if c else "" for c in rows[0]]
    _log(f"Columns found: {header}")

    def _col(keyword):
        for i, h in enumerate(header):
            if keyword.lower() in h.lower():
                return i
        return None

    col_date    = _col("Date")
    col_pv      = _col("Solar Production Energy")
    col_load    = _col("Load Consumed Energy")
    col_bat_chg = _col("Battery Charge Energy")
    col_bat_dis = _col("Battery Discharge Energy")
    col_imp     = _col("Grid Imported Energy")
    col_exp     = _col("Grid Exported Energy")

    if any(c is None for c in [col_date, col_pv, col_load, col_imp, col_exp]):
        _log("Could not map expected columns — check file format", level="ERROR")
        return 0

    init_daily_summary_db(db_path)
    imported = 0

    for row in rows[1:]:
        if not row or col_date >= len(row):
            continue
        raw_date = row[col_date]
        if raw_date is None or str(raw_date).strip().lower() in ("date", "total", ""):
            continue

        try:
            if isinstance(raw_date, str):
                d = date.fromisoformat(raw_date.strip())
            elif hasattr(raw_date, "date"):
                d = raw_date.date()
            else:
                continue
        except (ValueError, TypeError):
            continue

        def _num(col):
            if col is None or col >= len(row):
                return None
            v = row[col]
            if v is None:
                return None
            try:
                f = float(str(v).replace(",", "").strip())
                return f if f >= 0 else None
            except (ValueError, TypeError):
                return None

        pv      = _num(col_pv)
        home    = _num(col_load)
        imp     = _num(col_imp)
        exp     = _num(col_exp)
        bat_chg = _num(col_bat_chg)
        bat_dis = _num(col_bat_dis)

        if not any(v is not None and v > 0 for v in [pv, home, imp, exp]):
            continue

        row_dict = {
            "date":                     d.isoformat(),
            "pv_kwh":                   round(pv      or 0.0, 3),
            "home_kwh":                 round(home    or 0.0, 3),
            "grid_import_kwh":          round(imp     or 0.0, 3),
            "grid_export_kwh":          round(exp     or 0.0, 3),
            "battery_charge_kwh":       round(bat_chg or 0.0, 3),
            "battery_discharge_kwh":    round(bat_dis or 0.0, 3),
            "sigen_source":             "file_import",
            "elec_standing_charge_gbp": round(ELEC_STANDING_P_DAY / 100, 4),
            "gas_standing_charge_gbp":  round(GAS_STANDING_P_DAY  / 100, 4),
            "last_updated":             datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        }
        _upsert_sigen_row(db_path, row_dict)
        imported += 1

    _log(f"Imported {imported} days from Sigenergy File 2")
    return imported


def _upsert_sigen_row(db_path, row):
    """Upsert Sigenergy columns only — preserves Octopus fields already in the DB."""
    cols         = ", ".join(row.keys())
    placeholders = ", ".join(["?"] * len(row))
    sigen_cols   = [
        "pv_kwh", "home_kwh", "grid_import_kwh", "grid_export_kwh",
        "battery_charge_kwh", "battery_discharge_kwh", "sigen_source",
        "elec_standing_charge_gbp", "gas_standing_charge_gbp", "last_updated",
    ]
    update_clause = ", ".join(
        f"{c} = excluded.{c}" for c in sigen_cols if c in row
    )
    sql = (
        f"INSERT INTO daily_summary ({cols}) VALUES ({placeholders}) "
        f"ON CONFLICT(date) DO UPDATE SET {update_clause}"
    )
    con = sqlite3.connect(db_path)
    con.execute(sql, list(row.values()))
    con.commit()
    con.close()


# ---------------------------------------------------------------------------
# Sigenergy half-hourly aggregation
# ---------------------------------------------------------------------------

def _aggregate_halfhourly(timeseries_db, date_from, date_to, log_fn):
    """Aggregate halfhourly table rows into daily totals.

    Returns {date_str: {pv_kwh, home_kwh, imp_kwh, exp_kwh, bat_chg, bat_dis,
                        tracker_avg_p}}.
    Falls back gracefully if DB unavailable or table has no rows.
    """
    if not timeseries_db:
        return {}
    import os
    if not os.path.exists(timeseries_db):
        return {}

    result = {}
    try:
        con = sqlite3.connect(timeseries_db)
        rows = con.execute(
            """SELECT DATE(slot_start)                        AS day,
                      SUM(pv_kwh)                            AS pv,
                      SUM(home_kwh)                          AS home,
                      SUM(grid_import_kwh)                   AS imp,
                      SUM(grid_export_kwh)                   AS exp,
                      SUM(CASE WHEN battery_net_kwh > 0 THEN battery_net_kwh ELSE 0 END) AS bat_chg,
                      SUM(CASE WHEN battery_net_kwh < 0 THEN ABS(battery_net_kwh) ELSE 0 END) AS bat_dis,
                      SUM(grid_import_kwh * tracker_price_p) /
                          NULLIF(SUM(grid_import_kwh), 0)    AS tracker_avg_p
               FROM halfhourly
               WHERE slot_start >= ? AND slot_start < ?
               GROUP BY day""",
            (f"{date_from.isoformat()}T00:00:00",
             f"{(date_to + timedelta(days=1)).isoformat()}T00:00:00")
        ).fetchall()
        con.close()
        for r in rows:
            result[r[0]] = {
                "pv_kwh":       r[1] or 0.0,
                "home_kwh":     r[2] or 0.0,
                "imp_kwh":      r[3] or 0.0,
                "exp_kwh":      r[4] or 0.0,
                "bat_chg":      r[5] or 0.0,
                "bat_dis":      r[6] or 0.0,
                "tracker_avg_p": r[7],
            }
    except Exception as exc:
        log_fn(f"halfhourly aggregate failed: {exc}", level="WARNING")
    return result


# ---------------------------------------------------------------------------
# Existing daily_summary rows (for backfill-sourced data preservation)
# ---------------------------------------------------------------------------

def _load_existing_summary_rows(db_path, date_from, date_to):
    """Return {date_str: row_dict} for existing daily_summary rows in range."""
    try:
        con = sqlite3.connect(db_path)
        con.row_factory = sqlite3.Row
        rows = con.execute(
            "SELECT * FROM daily_summary WHERE date >= ? AND date <= ?",
            (date_from.isoformat(), date_to.isoformat())
        ).fetchall()
        con.close()
        return {r["date"]: dict(r) for r in rows}
    except Exception:
        return {}


# ---------------------------------------------------------------------------
# Octopus electricity consumption (import or export)
# ---------------------------------------------------------------------------

def _fetch_octopus_consumption(api_key, mpan, serial, direction, date_from, date_to, log_fn):
    """Fetch half-hourly electricity consumption from Octopus API.

    Returns {date_str: [(time_str, kwh), ...]} where time_str is local HH:MM.
    direction: 'electricity' for import, 'export' for export MPAN.
    """
    if not api_key or not mpan or not serial:
        return {}

    # Widen the UTC request window by an hour each side so the first local slots
    # (00:00 & 00:30) of the earliest day are not dropped in BST, when local
    # midnight is 23:00Z the previous day. Local-date bucketing below discards
    # any out-of-range days the widening pulls in.
    period_from = f"{(date_from - timedelta(days=1)).isoformat()}T23:00:00Z"
    period_to   = f"{(date_to + timedelta(days=1)).isoformat()}T01:00:00Z"

    url = (
        f"{_API_BASE}/electricity-meter-points/{mpan}/"
        f"meters/{serial}/consumption/"
    )
    params = {
        "period_from": period_from,
        "period_to":   period_to,
        "page_size":   _PAGE_SIZE,
        "order_by":    "period",
    }

    try:
        intervals = _paginate(url, api_key, params)
    except Exception as exc:
        log_fn(f"Octopus {direction} fetch failed: {exc}", level="WARNING")
        return {}

    log_fn(f"Octopus {direction}: {len(intervals)} half-hourly slots fetched")

    result = {}
    for item in intervals:
        try:
            start_utc = item.get("interval_start", "")
            kwh       = float(item.get("consumption", 0))
            if not start_utc or kwh < 0:
                continue
            local_dt  = _utc_to_local_dt(start_utc)
            ds        = local_dt.strftime("%Y-%m-%d")
            ts        = local_dt.strftime("%H:%M")
            result.setdefault(ds, []).append((ts, kwh))
        except (ValueError, TypeError):
            continue

    return result


# ---------------------------------------------------------------------------
# Octopus gas consumption
# ---------------------------------------------------------------------------

def _fetch_octopus_gas(api_key, mprn, gas_serial, date_from, date_to, log_fn):
    """Fetch daily gas consumption (m3) from Octopus API, converted to kWh.

    Returns {date_str: kwh}.
    """
    if not api_key or not mprn or not gas_serial:
        return {}

    # Widen the window an hour each side (see _fetch_octopus_consumption) so the
    # earliest local day's first slots aren't lost across the BST boundary.
    period_from = f"{(date_from - timedelta(days=1)).isoformat()}T23:00:00Z"
    period_to   = f"{(date_to + timedelta(days=1)).isoformat()}T01:00:00Z"

    url = (
        f"{_API_BASE}/gas-meter-points/{mprn}/"
        f"meters/{gas_serial}/consumption/"
    )
    params = {
        "period_from": period_from,
        "period_to":   period_to,
        "page_size":   _PAGE_SIZE,
        "order_by":    "period",
    }

    try:
        intervals = _paginate(url, api_key, params)
    except Exception as exc:
        log_fn(f"Octopus gas fetch failed: {exc}", level="WARNING")
        return {}

    log_fn(f"Octopus gas: {len(intervals)} intervals fetched")

    # Aggregate half-hourly m3 readings to daily kWh
    daily_m3 = {}
    for item in intervals:
        try:
            start_utc = item.get("interval_start", "")
            m3        = float(item.get("consumption", 0))
            if not start_utc or m3 < 0:
                continue
            local_dt  = _utc_to_local_dt(start_utc)
            ds        = local_dt.strftime("%Y-%m-%d")
            daily_m3[ds] = daily_m3.get(ds, 0.0) + m3
        except (ValueError, TypeError):
            continue

    # Convert m3 to kWh
    return {ds: round(m3 * GAS_KWH_PER_M3, 3) for ds, m3 in daily_m3.items()}


# ---------------------------------------------------------------------------
# Octopus Tracker historical rates
# ---------------------------------------------------------------------------

def _fetch_tracker_rates(date_from, date_to, region, log_fn):
    """Fetch historical Tracker (SILVER-25-04-11) unit rates for the date range.

    Returns {date_str: rate_p}. Uses public (unauthenticated) Octopus API.
    The Tracker tariff has one rate per day stored as standard-unit-rates entries.
    """
    tariff_code = f"{TRACKER_TARIFF_PREFIX}{region}"
    url = (
        f"{_API_BASE}/products/{TRACKER_PRODUCT}/electricity-tariffs/"
        f"{tariff_code}/standard-unit-rates/"
    )
    params = {
        "period_from": f"{date_from.isoformat()}T00:00:00Z",
        "period_to":   f"{(date_to + timedelta(days=1)).isoformat()}T00:00:00Z",
        "page_size":   500,
    }

    try:
        data = _api_get(url, params=params, api_key=None)
        results = data.get("results", [])
    except Exception as exc:
        log_fn(f"Tracker rate fetch failed: {exc}", level="WARNING")
        return {}

    rate_by_date = {}
    for item in results:
        try:
            valid_from = item.get("valid_from", "")
            rate_p     = float(item.get("value_inc_vat", 0))
            if not valid_from or rate_p <= 0:
                continue
            # valid_from is UTC midnight; Tracker rate applies to the UK calendar day
            local_dt = _utc_to_local_dt(valid_from)
            ds       = local_dt.strftime("%Y-%m-%d")
            rate_by_date[ds] = rate_p
        except (ValueError, TypeError):
            continue

    log_fn(f"Tracker rates: {len(rate_by_date)} days fetched")
    return rate_by_date


# ---------------------------------------------------------------------------
# Octopus Flux rates (live fetch)
# ---------------------------------------------------------------------------

def _fetch_flux_rates(region, log_fn):
    """Discover and return current Octopus Flux rates for the region.

    Returns dict: {offpeak_p, shoulder_p, peak_p}.
    Falls back to module-level constants if API unavailable.
    """
    fallback = {
        "offpeak_p":  FLUX_OFFPEAK_P,
        "shoulder_p": FLUX_SHOULDER_P,
        "peak_p":     FLUX_PEAK_P,
    }

    try:
        # Discover Flux import product code
        data     = _api_get(f"{_API_BASE}/products/", params={"page_size": 100})
        products = data.get("results", [])
        flux_code = None
        for p in products:
            if p.get("code", "").startswith("FLUX-IMPORT"):
                flux_code = p["code"]
                break
        if not flux_code:
            return fallback

        tariff_code = f"E-1R-{flux_code}-{region}"
        url = (
            f"{_API_BASE}/products/{flux_code}/electricity-tariffs/"
            f"{tariff_code}/standard-unit-rates/"
        )
        # Fetch today's Flux rate schedule
        today = date.today()
        params = {
            "period_from": f"{today.isoformat()}T00:00:00Z",
            "period_to":   f"{today.isoformat()}T23:59:59Z",
            "page_size":   50,
        }
        data    = _api_get(url, params=params)
        results = data.get("results", [])

        offpeak_rates  = []
        shoulder_rates = []
        peak_rates     = []

        for item in results:
            valid_from = item.get("valid_from", "")
            rate_p     = float(item.get("value_inc_vat", 0))
            if not valid_from:
                continue
            local_dt = _utc_to_local_dt(valid_from)
            hhmm     = local_dt.strftime("%H:%M")

            if _in_window(hhmm, FLUX_OFFPEAK_START, FLUX_OFFPEAK_END):
                offpeak_rates.append(rate_p)
            elif _in_window(hhmm, FLUX_PEAK_START, FLUX_PEAK_END):
                peak_rates.append(rate_p)
            else:
                shoulder_rates.append(rate_p)

        result = {
            "offpeak_p":  round(sum(offpeak_rates)  / len(offpeak_rates),  4) if offpeak_rates  else FLUX_OFFPEAK_P,
            "shoulder_p": round(sum(shoulder_rates) / len(shoulder_rates), 4) if shoulder_rates else FLUX_SHOULDER_P,
            "peak_p":     round(sum(peak_rates)     / len(peak_rates),     4) if peak_rates     else FLUX_PEAK_P,
        }
        log_fn(f"Flux rates fetched: off-peak={result['offpeak_p']}p  "
               f"shoulder={result['shoulder_p']}p  peak={result['peak_p']}p")
        return result

    except Exception as exc:
        log_fn(f"Flux rate fetch failed (using defaults): {exc}", level="WARNING")
        return fallback


# ---------------------------------------------------------------------------
# TOU cost calculators
# ---------------------------------------------------------------------------

def _apply_tou_simple(slots, cheap_start, cheap_end, cheap_p, peak_p):
    """Apply a simple two-band TOU rate to a list of (hhmm_str, kwh) tuples."""
    total_p = 0.0
    for hhmm, kwh in slots:
        rate = cheap_p if _in_window(hhmm, cheap_start, cheap_end) else peak_p
        total_p += kwh * rate
    return total_p


def _apply_tou_flux(slots, offpeak_p, shoulder_p, peak_p):
    """Apply Flux three-band TOU rate to a list of (hhmm_str, kwh) tuples."""
    total_p = 0.0
    for hhmm, kwh in slots:
        if _in_window(hhmm, FLUX_OFFPEAK_START, FLUX_OFFPEAK_END):
            rate = offpeak_p
        elif _in_window(hhmm, FLUX_PEAK_START, FLUX_PEAK_END):
            rate = peak_p
        else:
            rate = shoulder_p
        total_p += kwh * rate
    return total_p


def _in_window(hhmm, start, end):
    """True if hhmm falls in [start, end) — handles overnight windows."""
    if start <= end:
        return start <= hhmm < end
    return hhmm >= start or hhmm < end


# ---------------------------------------------------------------------------
# DB upsert
# ---------------------------------------------------------------------------

def _upsert_row(db_path, row):
    """INSERT OR REPLACE one row into daily_summary."""
    cols = ", ".join(row.keys())
    placeholders = ", ".join(["?"] * len(row))
    sql  = f"INSERT OR REPLACE INTO daily_summary ({cols}) VALUES ({placeholders})"
    con  = sqlite3.connect(db_path)
    con.execute(sql, list(row.values()))
    con.commit()
    con.close()


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _api_get(url, params=None, api_key=None):
    """GET Octopus API endpoint. Returns parsed JSON."""
    if params:
        url = url + "?" + urllib.parse.urlencode(params)

    headers = {
        "Accept":     "application/json",
        "User-Agent": "IndigoTariffAnalyser/1.0",
    }
    if api_key:
        creds = base64.b64encode(f"{api_key}:".encode()).decode()
        headers["Authorization"] = f"Basic {creds}"

    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, timeout=_REQUEST_TIMEOUT) as resp:
        return json.loads(resp.read().decode("utf-8"))


def _paginate(url, api_key, params):
    """Follow Octopus API pagination, returning all results."""
    all_results = []
    next_url    = url
    next_params = params
    while next_url:
        data = _api_get(next_url, params=next_params, api_key=api_key)
        all_results.extend(data.get("results", []))
        next_url    = data.get("next")
        next_params = None
        if len(all_results) > 100000:
            break
    return all_results


def _utc_to_local_dt(utc_str):
    """Parse a UTC ISO timestamp and return a local datetime.

    Handles all formats Octopus API returns:
      "2026-01-04T00:00:00Z"              (Tracker rates)
      "2026-01-01T00:00:00+00:00"         (some endpoints)
      "2026-01-01T00:00:00.000000+00:00"  (consumption — has sub-seconds)
    """
    s = utc_str.split(".")[0].rstrip("Z")   # strip sub-seconds and trailing Z
    if "+" in s[10:]:
        s = s[:19]                           # strip +HH:MM offset if present
    dt_utc = datetime.fromisoformat(s + "+00:00")
    return dt_utc.astimezone(tz=None)
